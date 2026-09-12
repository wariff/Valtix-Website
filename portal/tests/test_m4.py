#!/usr/bin/env python3
"""Tests fuer M4: Zuordnung, Pruefliste, Freigabe, Ausgabe."""
import os
import sys
import tempfile

import pytest

HIER = os.path.dirname(os.path.abspath(__file__))
PORTAL = os.path.dirname(HIER)
sys.path.insert(0, PORTAL)
sys.path.insert(0, os.path.join(HIER, 'beispiele'))

import erzeugen                                            # noqa: E402
import mapping as mp                                       # noqa: E402

XLSX_MIME = ('application/vnd.openxmlformats-officedocument'
             '.spreadsheetml.sheet')


@pytest.fixture()
def umgebung():
    os.environ['VALTIX_DB'] = tempfile.mktemp(suffix='.sqlite3')
    os.environ['VALTIX_ABLAGE'] = tempfile.mkdtemp()
    for name in ('datenbank', 'speicher', 'perioden', 'benachrichtigung',
                 'aufgaben', 'mapping', 'uebernahme'):
        sys.modules.pop(name, None)
    import datenbank, perioden, aufgaben, uebernahme, mapping   # noqa: E402
    datenbank.anlegen()
    mid = datenbank.mandant_anlegen('Alpha GmbH')
    perioden.checkliste_uebernehmen(mid)
    yield {'db': datenbank, 'pd': perioden, 'auf': aufgaben, 'ue': uebernahme,
           'mp': mapping, 'mid': mid}


def _gefuellt(u, jahr_monat='2026-07'):
    pd, auf, mid = u['pd'], u['auf'], u['mid']
    pd.dokument_ablegen(mid, jahr_monat, 'susa', 'susa.csv', 'text/csv',
                        erzeugen.susa_csv(), None)
    pd.dokument_ablegen(mid, jahr_monat, 'bwa', 'bwa.xlsx', XLSX_MIME,
                        erzeugen.bwa_xlsx(), None)
    auf.abarbeiten()
    return pd.periode(mid, jahr_monat)


# ------------------------------------------------------------------ Mapping
def test_kontenrahmen_wird_erkannt():
    assert mp.rahmen_erkennen([8400, 3200, 4120]) == 'SKR03'
    assert mp.rahmen_erkennen([4000, 5200, 6000]) == 'SKR04'
    assert mp.rahmen_erkennen([]) == 'unbekannt'


def test_konto_schlaegt_bezeichnung():
    ziel, konfidenz, woher = mp.vorschlag('8400', 'irgendwas', 'SKR03')
    assert ziel == 'hauptleistung' and woher.startswith('Kontenbereich')
    assert konfidenz > 0.7


def test_bezeichnung_greift_ohne_konto():
    ziel, _, woher = mp.vorschlag('', 'Personalaufwand inkl. AG-Anteil', 'SKR04')
    assert ziel == 'personal' and woher == 'Bezeichnung'


def test_unbekanntes_bleibt_offen():
    assert mp.vorschlag('9999', 'Kreatives Konto', 'SKR04')[0] is None


def test_gelernte_regel_schlaegt_alles(umgebung):
    mp2, mid = umgebung['mp'], umgebung['mid']
    mp2.regel_merken(mid, '8400', 'Erlöse 19%', 'warenverkauf')
    gelernt = mp2.gelernte_regeln(mid)
    ziel, konfidenz, woher = mp2.vorschlag('8400', 'Erlöse 19%', 'SKR03', gelernt)
    assert (ziel, konfidenz, woher) == ('warenverkauf', 1.0, 'gelernt')


def test_alle_zielfelder_zeigen_auf_die_vorlage():
    import openpyxl
    vorlage = os.path.join(os.path.dirname(PORTAL), 'tools', 'bericht',
                           'VALTIX_Eingabevorlage.xlsx')
    wb = openpyxl.load_workbook(vorlage)
    for schluessel, (blatt, zeile, klartext, _) in mp.ZIELFELDER.items():
        assert blatt in wb.sheetnames, schluessel
        # In Spalte A der Zeile muss eine Beschriftung stehen, keine leere Zelle.
        assert wb[blatt].cell(zeile, 1).value, f'{schluessel} -> {blatt} {zeile}'


# ---------------------------------------------------------------- Uebernahme
def test_werte_werden_nicht_doppelt_gezaehlt(umgebung):
    ue = umgebung['ue']
    p = _gefuellt(umgebung)
    liste = ue.pruefliste(p['id'])
    umsatz = [z for z in liste['zeilen'] if z['schluessel'] == 'hauptleistung'][0]
    # BWA und Summenliste nennen beide 365.057,97. Die Summe waere doppelt.
    assert round(umsatz['wert'], 2) == 365057.97
    personal = [z for z in liste['zeilen'] if z['schluessel'] == 'personal'][0]
    assert round(personal['wert'], 2) == 142062.35


def test_offener_posten_landet_in_der_klaerliste(umgebung):
    ue = umgebung['ue']
    p = _gefuellt(umgebung)
    liste = ue.pruefliste(p['id'])
    assert any('Betriebsergebnis' in (k['quelle'] or '') for k in liste['klaerliste'])


def test_klaerfall_zuordnen_lernt_die_regel(umgebung):
    ue, mp2, mid = umgebung['ue'], umgebung['mp'], umgebung['mid']
    p = _gefuellt(umgebung)
    liste = ue.pruefliste(p['id'])
    fall = liste['klaerliste'][0]
    ue.klaerfall_zuordnen(p['id'], fall['id'], 'sonstige_ertrag', None)
    assert mp2.gelernte_regeln(mid).get(fall['quelle']) == 'sonstige_ertrag'
    assert not ue.pruefliste(p['id'])['klaerliste']


def test_pflichtfeld_ohne_wert_blockiert_die_freigabe(umgebung):
    ue, pd, mid, auf = umgebung['ue'], umgebung['pd'], umgebung['mid'], umgebung['auf']
    pd.dokument_ablegen(mid, '2026-07', 'susa', 'nur-bank.csv', 'text/csv',
                        b'Konto;Bezeichnung;Saldo\r\n1200;Bank;121.560,45', None)
    auf.abarbeiten()
    p = pd.periode(mid, '2026-07')
    with pytest.raises(pd.Verweigert) as e:
        ue.freigeben(p['id'], None)
    assert 'rechnet der Bericht nicht' in str(e.value)


def test_freigabe_setzt_status_und_werte(umgebung):
    ue, pd = umgebung['ue'], umgebung['pd']
    p = _gefuellt(umgebung)
    anzahl = ue.freigeben(p['id'], None)
    assert anzahl >= 3
    assert pd.periode_nach_id(p['id'])['status'] == 'freigegeben'
    w = ue.werte(p['id'])
    assert round(w['hauptleistung']['wert'], 2) == 365057.97
    assert w['hauptleistung']['freigegeben_am']


def test_handkorrektur_wird_protokolliert(umgebung):
    ue, db = umgebung['ue'], umgebung['db']
    p = _gefuellt(umgebung)
    ue.freigeben(p['id'], None)
    ue.wert_setzen(p['id'], 'hauptleistung', 400000.0, None)
    eintrag = [e for e in db.protokoll() if e['ereignis'] == 'wert_freigegeben'][0]
    assert eintrag['vorher'] == '365057.97' and eintrag['nachher'] == '400000.00'


def test_grosse_abweichung_zum_vormonat_wird_gemeldet(umgebung):
    ue, pd, mid, auf = umgebung['ue'], umgebung['pd'], umgebung['mid'], umgebung['auf']
    vor = _gefuellt(umgebung, '2026-06')
    ue.freigeben(vor['id'], None)
    pd.dokument_ablegen(mid, '2026-07', 'susa', 'juli.csv', 'text/csv',
                        b'Konto;Bezeichnung;Saldo\r\n4000;Umsatzerl\xf6se;900.000,00\r\n'
                        b'5000;Wareneinsatz;100.000,00\r\n6000;L\xf6hne;50.000,00',
                        None)
    auf.abarbeiten()
    p = pd.periode(mid, '2026-07')
    liste = ue.pruefliste(p['id'])
    umsatz = [z for z in liste['zeilen'] if z['schluessel'] == 'hauptleistung'][0]
    assert any('Vormonat' in w for w in umsatz['warnungen']), umsatz['warnungen']


def test_negativer_wert_wird_gemeldet(umgebung):
    ue = umgebung['ue']
    p = _gefuellt(umgebung)
    ue.wert_setzen(p['id'], 'hauptleistung', -100.0, None)
    liste = ue.pruefliste(p['id'])
    umsatz = [z for z in liste['zeilen'] if z['schluessel'] == 'hauptleistung'][0]
    assert any('Vorzeichen' in w for w in umsatz['warnungen'])


# ------------------------------------------------------------------- Ausgabe
def test_export_json_nennt_blatt_und_zeile(umgebung):
    import json
    ue = umgebung['ue']
    p = _gefuellt(umgebung)
    ue.freigeben(p['id'], None)
    d = json.loads(ue.als_json(p['id']))
    assert d['jahr_monat'] == '2026-07'
    assert d['felder']['hauptleistung']['blatt'] == '2 GuV'
    assert d['felder']['hauptleistung']['zeile'] == 6


def test_export_xlsx_schreibt_in_die_monatsspalte(umgebung):
    import io
    import openpyxl
    ue = umgebung['ue']
    p = _gefuellt(umgebung)
    ue.freigeben(p['id'], None)
    wb = openpyxl.load_workbook(io.BytesIO(ue.als_xlsx(p['id'])))
    # Juli ist die achte Spalte, H
    assert round(wb['2 GuV']['H6'].value, 2) == 365057.97
    assert round(wb['2 GuV']['H27'].value, 2) == 142062.35
    assert wb['1 Stammdaten']['B9'].value == 'Juli'


def test_formeln_in_zellen_werden_entschaerft():
    from uebernahme import _sicher_fuer_excel
    assert _sicher_fuer_excel('=1+1') == "'=1+1"
    assert _sicher_fuer_excel('-SUMME()') == "'-SUMME()"
    assert _sicher_fuer_excel('Miete') == 'Miete'
    assert _sicher_fuer_excel(1234.5) == 1234.5
