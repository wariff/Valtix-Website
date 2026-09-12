#!/usr/bin/env python3
"""Aus dem Gelesenen einen Vorschlag machen, pruefen, freigeben, ausgeben.

Kein Wert wird automatisch produktiv. Der Ablauf ist immer: vorschlagen,
in der Pruefansicht ansehen, korrigieren, freigeben. Korrekturen lernt das
Portal je Mandant und benutzt sie im Folgemonat wieder.
"""
import io
import json
import os
import re
import sys

HIER = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HIER)
import datenbank as db                                    # noqa: E402
import leser                                              # noqa: E402
import mapping as mp                                      # noqa: E402
import perioden as pd                                     # noqa: E402

KONTO_SPALTEN = ('konto', 'kto', 'kontonr', 'kontonummer', 'sachkonto')
TEXT_SPALTEN = ('bezeichnung', 'text', 'position', 'kontenbezeichnung', 'name')
BETRAG_SPALTEN = ('saldo', 'betrag', 'monat', 'wert', 'summe', 'ist')
# Pflichtfelder, ohne die der Bericht nicht rechnen kann.
PFLICHT = ('hauptleistung', 'wareneinsatz', 'personal')
ABWEICHUNG = 0.30          # ab 30 Prozent Unterschied zum Vormonat nachfragen


def _kopfzeile(zeilen):
    """Findet die Zeile mit den Spaltennamen und die Spaltenpositionen."""
    for nr, zeile in enumerate(zeilen[:6]):
        klein = [str(z or '').strip().lower() for z in zeile]
        konto = next((i for i, z in enumerate(klein) if z in KONTO_SPALTEN), None)
        text = next((i for i, z in enumerate(klein) if z in TEXT_SPALTEN), None)
        betrag = next((i for i, z in enumerate(klein) if z in BETRAG_SPALTEN), None)
        if betrag is not None and (konto is not None or text is not None):
            return nr, konto, text, betrag
    return None, None, None, None


def posten(tabelle):
    """Gibt (quelle, bezeichnung, betrag) je Zeile, so gut es geht."""
    zeilen = tabelle.get('zeilen', [])
    kopf, i_konto, i_text, i_betrag = _kopfzeile(zeilen)
    raus = []
    if kopf is None:
        # Ohne Kopfzeile: erste Zelle als Bezeichnung, letzte Zahl als Betrag.
        for z in zeilen:
            zahlen = [(i, leser._zahl(w)) for i, w in enumerate(z)]
            zahlen = [(i, v) for i, v in zahlen if v is not None]
            if not zahlen or not str(z[0] or '').strip():
                continue
            raus.append((str(z[0]).strip(), str(z[0]).strip(), zahlen[-1][1]))
        return raus
    for z in zeilen[kopf + 1:]:
        if i_betrag >= len(z):
            continue
        betrag = leser._zahl(z[i_betrag])
        if betrag is None:
            continue
        quelle = str(z[i_konto]).strip() if i_konto is not None and i_konto < len(z) else ''
        text = str(z[i_text]).strip() if i_text is not None and i_text < len(z) else ''
        if not quelle and not text:
            continue
        raus.append((quelle or text, text or quelle, betrag))
    return raus


def vorschlagen(periode_id):
    """Geht alle Extraktionen der Periode durch und schlaegt Werte vor."""
    p = pd.periode_nach_id(periode_id)
    if not p:
        raise pd.Verweigert('Diese Periode gibt es nicht.')
    gelernt = mp.gelernte_regeln(p['mandant_id'])
    with db.verbinden() as con:
        dokumente = [dict(r) for r in con.execute(
            'SELECT d.id, d.dateiname, e.tabellen, e.weg FROM dokument d '
            'JOIN extraktion e ON e.dokument_id=d.id '
            "WHERE d.periode_id=? AND d.aktiv=1 AND e.status='roh'",
            (periode_id,)).fetchall()]

    alle, hinweise = [], []
    for d in dokumente:
        if d['weg'] == 'datev':
            # Ein Buchungsstapel listet einzelne Buchungen, keine Periodensummen.
            # Die Spaltenfolge ist eine andere als in BWA und Saldenliste, eine
            # Summe daraus waere geraten. Deshalb bleibt die Datei hier liegen.
            hinweise.append({'datei': d['dateiname'],
                             'text': 'Buchungsstapel, daraus werden keine Werte '
                                     'vorgeschlagen. Bitte BWA oder Summen- und '
                                     'Saldenliste hochladen.'})
            continue
        for tabelle in json.loads(d['tabellen'] or '[]'):
            for quelle, text, betrag in posten(tabelle):
                alle.append({'dokument_id': d['id'], 'datei': d['dateiname'],
                             'quelle': quelle, 'bezeichnung': text, 'betrag': betrag})

    konten = [k for k in (mp.kontonummer(a['quelle']) for a in alle) if k]
    rahmen = mp.rahmen_erkennen(konten)

    # Innerhalb einer Datei werden mehrere Konten zu einem Feld addiert.
    # Ueber Dateien hinweg NICHT: BWA und Summen- und Saldenliste enthalten
    # dieselben Zahlen, das haette jeden Betrag verdoppelt. Stattdessen wird
    # je Feld die zuverlaessigste Datei genommen und eine Abweichung gemeldet.
    je_datei, offen = {}, []
    for a in alle:
        ziel, konfidenz, woher = mp.vorschlag(a['quelle'], a['bezeichnung'],
                                              rahmen, gelernt)
        if not ziel:
            offen.append(a)
            continue
        f = je_datei.setdefault((a['dokument_id'], ziel),
                                {'wert': 0.0, 'konfidenz': 1.0, 'posten': [],
                                 'dokument_id': a['dokument_id'],
                                 'datei': a['datei']})
        f['wert'] += abs(a['betrag'])
        f['konfidenz'] = min(f['konfidenz'], konfidenz)
        f['posten'].append({**a, 'woher': woher})

    felder = {}
    for (dokument_id, ziel), f in je_datei.items():
        felder.setdefault(ziel, []).append(f)
    for ziel, kandidaten in list(felder.items()):
        # Beste Quelle: hoechste Konfidenz, bei Gleichstand die mit mehr Posten.
        kandidaten.sort(key=lambda k: (k['konfidenz'], len(k['posten'])), reverse=True)
        beste = dict(kandidaten[0])
        beste['abweichungen'] = []
        for andere in kandidaten[1:]:
            if beste['wert'] and abs(andere['wert'] - beste['wert']) / abs(beste['wert']) > 0.01:
                beste['abweichungen'].append(
                    {'datei': andere['datei'], 'wert': andere['wert']})
        felder[ziel] = beste

    with db.verbinden() as con:
        con.execute('DELETE FROM klaerfall WHERE periode_id=? AND erledigt=0',
                    (periode_id,))
        for a in offen:
            con.execute('INSERT INTO klaerfall (periode_id, dokument_id, quelle, '
                        'bezeichnung, betrag, erstellt_am) VALUES (?,?,?,?,?,?)',
                        (periode_id, a['dokument_id'], a['quelle'],
                         a['bezeichnung'], a['betrag'], db.jetzt()))
    return {'rahmen': rahmen, 'felder': felder, 'offen': offen,
            'hinweise': hinweise, 'posten_gesamt': len(alle)}


def _vormonat(p):
    jahr, monat = int(p['jahr_monat'][:4]), int(p['jahr_monat'][5:])
    jahr, monat = (jahr, monat - 1) if monat > 1 else (jahr - 1, 12)
    return pd.periode(p['mandant_id'], f'{jahr}-{monat:02d}')


def werte(periode_id):
    with db.verbinden() as con:
        return {r['feldschluessel']: dict(r) for r in con.execute(
            'SELECT * FROM kennzahl_wert WHERE periode_id=?', (periode_id,)).fetchall()}


def pruefliste(periode_id):
    """Was in der Pruefansicht steht: Wert, Herkunft, Konfidenz, Auffaelliges."""
    p = pd.periode_nach_id(periode_id)
    vorschlag = vorschlagen(periode_id)
    schon = werte(periode_id)
    vor = _vormonat(p)
    vorwerte = werte(vor['id']) if vor else {}

    zeilen = []
    for schluessel, (blatt, zeile, klartext, _) in mp.ZIELFELDER.items():
        v = vorschlag['felder'].get(schluessel)
        fest = schon.get(schluessel)
        wert = fest['wert'] if fest else (v['wert'] if v else None)
        warnungen = []
        if wert is not None and wert < 0:
            warnungen.append('negativer Betrag, Vorzeichen prüfen')
        alt = vorwerte.get(schluessel, {}).get('wert')
        if wert is not None and alt:
            unterschied = abs(wert - alt) / abs(alt)
            if unterschied > ABWEICHUNG:
                warnungen.append(f'{unterschied * 100:.0f} % Unterschied zum Vormonat')
        if schluessel in PFLICHT and not wert:
            warnungen.append('Pflichtfeld ohne Wert')
        for ab in (v or {}).get('abweichungen', []):
            warnungen.append(f'{ab["datei"]} nennt {ab["wert"]:,.2f}'
                             .replace(',', '.') + ' statt dessen')
        zeilen.append({
            'schluessel': schluessel, 'klartext': klartext, 'blatt': blatt,
            'zeile': zeile, 'wert': wert,
            'konfidenz': (fest['konfidenz'] if fest else (v['konfidenz'] if v else None)),
            'posten': v['posten'] if v else [],
            'freigegeben': bool(fest and fest['freigegeben_am']),
            'warnungen': warnungen,
        })

    aktiva = next((z['wert'] for z in zeilen if z['schluessel'] == 'bilanzsumme'), None)
    ek = next((z['wert'] for z in zeilen if z['schluessel'] == 'eigenkapital'), None)
    if aktiva and ek and ek > aktiva:
        for z in zeilen:
            if z['schluessel'] == 'bilanzsumme':
                z['warnungen'].append('Eigenkapital größer als die Bilanzsumme')

    with db.verbinden() as con:
        offen = [dict(r) for r in con.execute(
            'SELECT * FROM klaerfall WHERE periode_id=? AND erledigt=0 ORDER BY id',
            (periode_id,)).fetchall()]
    return {'rahmen': vorschlag['rahmen'], 'zeilen': zeilen, 'klaerliste': offen,
            'hinweise': vorschlag.get('hinweise', []),
            'posten_gesamt': vorschlag['posten_gesamt']}


def wert_setzen(periode_id, schluessel, wert, von, herkunft='von Hand'):
    if schluessel not in mp.ZIELFELDER:
        raise pd.Verweigert('Unbekanntes Feld.')
    alt = werte(periode_id).get(schluessel, {}).get('wert')
    with db.verbinden() as con:
        con.execute('INSERT INTO kennzahl_wert (periode_id, feldschluessel, wert, '
                    'herkunft, konfidenz, freigegeben_von, freigegeben_am) '
                    'VALUES (?,?,?,?,?,?,?) '
                    'ON CONFLICT(periode_id, feldschluessel) DO UPDATE SET '
                    'wert=excluded.wert, herkunft=excluded.herkunft, '
                    'freigegeben_von=excluded.freigegeben_von, '
                    'freigegeben_am=excluded.freigegeben_am',
                    (periode_id, schluessel, wert, herkunft, 1.0, von, db.jetzt()))
    db.protokollieren('wert_freigegeben', benutzer_id=von,
                      detail=f'periode {periode_id}, {schluessel}',
                      vorher=None if alt is None else f'{alt:.2f}',
                      nachher=None if wert is None else f'{wert:.2f}')


def freigeben(periode_id, von):
    """Uebernimmt alle vorgeschlagenen Werte, die noch nicht gesetzt sind."""
    liste = pruefliste(periode_id)
    fehlend = [z['klartext'] for z in liste['zeilen']
               if z['schluessel'] in PFLICHT and not z['wert']]
    if fehlend:
        raise pd.Verweigert('Ohne diese Felder rechnet der Bericht nicht: '
                            + ', '.join(fehlend))
    for z in liste['zeilen']:
        if z['wert'] is not None and not z['freigegeben']:
            wert_setzen(periode_id, z['schluessel'], z['wert'], von,
                        herkunft='Vorschlag übernommen')
    pd.status_setzen(periode_id, 'freigegeben', von)
    return sum(1 for z in liste['zeilen'] if z['wert'] is not None)


def klaerfall_zuordnen(periode_id, klaerfall_id, zielfeld, von):
    """Ordnet einen offenen Posten zu und merkt sich die Regel."""
    p = pd.periode_nach_id(periode_id)
    with db.verbinden() as con:
        k = con.execute('SELECT * FROM klaerfall WHERE id=? AND periode_id=?',
                        (klaerfall_id, periode_id)).fetchone()
    if not k:
        raise pd.Verweigert('Dieser Klärfall gehört nicht zu dieser Periode.')
    mp.regel_merken(p['mandant_id'], k['quelle'], k['bezeichnung'], zielfeld, von)
    with db.verbinden() as con:
        con.execute('UPDATE klaerfall SET erledigt=1 WHERE id=?', (klaerfall_id,))


# ------------------------------------------------------------------- Ausgabe
def _sicher_fuer_excel(wert):
    """Excel liest Zellen, die mit = + - @ beginnen, als Formel. Ein
    vorangestelltes Apostroph entschaerft das."""
    if isinstance(wert, str) and wert[:1] in ('=', '+', '-', '@'):
        return "'" + wert
    return wert


def als_json(periode_id):
    p = pd.periode_nach_id(periode_id)
    w = werte(periode_id)
    return json.dumps({
        'jahr_monat': p['jahr_monat'],
        'felder': {k: {'wert': v['wert'], 'blatt': mp.ZIELFELDER[k][0],
                       'zeile': mp.ZIELFELDER[k][1], 'herkunft': v['herkunft']}
                   for k, v in w.items() if k in mp.ZIELFELDER},
    }, ensure_ascii=False, indent=2)


def als_xlsx(periode_id, vorlage=None):
    """Schreibt die freigegebenen Werte in eine Kopie der Eingabevorlage,
    in die Spalte des Berichtsmonats."""
    import openpyxl
    vorlage = vorlage or os.path.join(os.path.dirname(HIER), 'tools', 'bericht',
                                      'VALTIX_Eingabevorlage.xlsx')
    p = pd.periode_nach_id(periode_id)
    monat = int(p['jahr_monat'][5:])
    spalte = 1 + monat                      # Spalte B ist Januar
    wb = openpyxl.load_workbook(vorlage)
    for schluessel, v in werte(periode_id).items():
        if schluessel not in mp.ZIELFELDER:
            continue
        blatt, zeile, _, vorzeichen = mp.ZIELFELDER[schluessel]
        if blatt not in wb.sheetnames:
            continue
        wb[blatt].cell(zeile, spalte).value = _sicher_fuer_excel(
            None if v['wert'] is None else v['wert'] * vorzeichen)
    name = wb.sheetnames[1] if len(wb.sheetnames) > 1 else wb.sheetnames[0]
    if name == '1 Stammdaten':
        wb[name]['B8'] = int(p['jahr_monat'][:4])
        wb[name]['B9'] = pd.MONATE[monat - 1]
    puffer = io.BytesIO()
    wb.save(puffer)
    return puffer.getvalue()
