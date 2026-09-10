#!/usr/bin/env python3
"""Liest die Eingabevorlage als Monatsraster.

modell.py rechnet einen einzelnen Berichtsmonat durch. Hier geht es um die
Breite: jede Position ueber alle befuellten Monate, so wie sie im Portal
nebeneinander stehen soll.

Auch hier gilt: gelesen werden nur Eingabefelder. Die Summenzeilen der Vorlage
sind Excel-Formeln, deren Ergebnis in einer nie geoeffneten Datei fehlt.
"""
import os
import sys
import warnings

import openpyxl

warnings.filterwarnings('ignore')
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from modell import (MONATE, GUV_ERLOESE, GUV_SONSTIGE, GUV_VARIABEL, GUV_FIX,
                    GUV_AFA, GUV_ZINSAUFWAND, GUV_ZINSERTRAG, MEN, LIQ,
                    ZIEL_ZEILEN, EIGENE_ZEILEN)

# Spalte B ist Januar
SPALTE = {m: 2 + i for i, m in enumerate(MONATE)}


def _z(w):
    return 0.0 if w is None or isinstance(w, str) else float(w)


def _w(w):
    return None if w is None or isinstance(w, str) else float(w)


def _teile(a, b):
    """Division, die bei fehlendem Nenner nichts behauptet."""
    if a is None or not b:
        return None
    return a / b


class Zeile:
    def __init__(self, schluessel, titel, art, werte, ebene=0, gruppe=None,
                 vorzeichen=1):
        self.schluessel = schluessel
        self.titel = titel
        self.art = art              # geld, prozent, zahl, tage
        self.werte = werte          # Liste in der Reihenfolge der Monate
        self.ebene = ebene          # 0 Summe, 1 Gruppe, 2 Einzelposition
        self.gruppe = gruppe        # Schluessel der Gruppe, zu der die Zeile klappt
        self.vorzeichen = vorzeichen
        self.ziel = None
        self.richtung = None

    @property
    def gefuellt(self):
        return any(w for w in self.werte)


class Matrix:
    """Alle Positionen der Vorlage ueber alle befuellten Monate."""

    def __init__(self, pfad):
        wb = openpyxl.load_workbook(pfad, data_only=False)
        self.stamm = wb.worksheets[1]
        self.guv = wb.worksheets[2]
        self.men = wb.worksheets[3]
        self.liq = wb.worksheets[4]
        self.ziel = wb.worksheets[5]

        self.firma = self.stamm['B5'].value or ''
        self.branche = self.stamm['B7'].value or ''
        self.jahr = self.stamm['B8'].value or ''
        self.berichtsmonat = (self.stamm['B9'].value or '').strip()
        self.einheit = self.stamm['B14'].value or 'Leistungseinheit'
        self.einheit_plural = self.stamm['B15'].value or self.einheit

        self.monate = self._befuellte_monate()
        self.aktiv = (self.monate.index(self.berichtsmonat)
                      if self.berichtsmonat in self.monate else len(self.monate) - 1)

        self.bloecke = []
        self._ergebnis()
        self._operativ()
        self._liquiditaet()
        self._ziele_zuordnen()

    # ------------------------------------------------------------- Grundlagen
    def _befuellte_monate(self):
        """Ein Monat zaehlt, sobald in der GuV irgendein Erloes eingetragen ist."""
        raus = []
        for m in MONATE:
            sp = SPALTE[m]
            if any(_z(self.guv.cell(z, sp).value) for z in GUV_ERLOESE.values()):
                raus.append(m)
        return raus

    def _reihe(self, blatt, zeile, umrechnung=_z):
        return [umrechnung(blatt.cell(zeile, SPALTE[m]).value) for m in self.monate]

    def _titel(self, blatt, zeile, ersatz):
        """Beschriftung aus der Datei. Mandanten benennen Zeilen um, dann steht
        im Portal ihr Wort und nicht meines."""
        t = blatt.cell(zeile, 1).value
        t = str(t).strip() if t else ''
        if not t:
            return ersatz
        return t.replace(' (optional)', '').strip()

    @staticmethod
    def _summe(*reihen):
        return [sum(w[i] for w in reihen) for i in range(len(reihen[0]))]

    def _gruppe(self, blatt, zeilen, schluessel, titel, art='geld'):
        """Summenzeile plus die Einzelpositionen, die tatsaechlich Werte haben."""
        kinder = []
        for ersatz, z in zeilen.items():
            werte = self._reihe(blatt, z)
            if any(werte):
                kinder.append(Zeile(f'{schluessel}-{z}', self._titel(blatt, z, ersatz),
                                    art, werte, ebene=2, gruppe=schluessel))
        summe = [sum(k.werte[i] for k in kinder) for i in range(len(self.monate))] \
            if kinder else [0.0] * len(self.monate)
        kopf = Zeile(schluessel, titel, art, summe, ebene=1)
        return kopf, kinder

    # ---------------------------------------------------------------- Bloecke
    def _ergebnis(self):
        z = []
        umsatz_kopf, umsatz_kinder = self._gruppe(
            self.guv, GUV_ERLOESE, 'umsatz', 'Umsatzerlöse')
        z.append(umsatz_kopf)
        z.extend(umsatz_kinder)

        sonstige = self._reihe(self.guv, GUV_SONSTIGE)
        if any(sonstige):
            z.append(Zeile('sonstige_ertrag',
                           self._titel(self.guv, GUV_SONSTIGE, 'Sonstige betriebliche Erträge'),
                           'geld', sonstige, ebene=1))
        gesamt = self._summe(umsatz_kopf.werte, sonstige)
        z.append(Zeile('gesamtleistung', 'Gesamtleistung', 'geld', gesamt))

        var_kopf, var_kinder = self._gruppe(
            self.guv, GUV_VARIABEL, 'variabel', 'Variable Kosten', 'geld')
        var_kopf.vorzeichen = -1
        for k in var_kinder:
            k.vorzeichen = -1
        z.append(var_kopf)
        z.extend(var_kinder)

        db = [gesamt[i] - var_kopf.werte[i] for i in range(len(gesamt))]
        z.append(Zeile('db', 'Deckungsbeitrag', 'geld', db))
        z.append(Zeile('db_quote', 'Deckungsbeitragsquote', 'prozent',
                       [_teile(db[i], gesamt[i]) and db[i] / gesamt[i] * 100
                        if gesamt[i] else None for i in range(len(gesamt))], ebene=2))

        fix_kopf, fix_kinder = self._gruppe(
            self.guv, GUV_FIX, 'fix', 'Fixkosten', 'geld')
        fix_kopf.vorzeichen = -1
        for k in fix_kinder:
            k.vorzeichen = -1
        z.append(fix_kopf)
        z.extend(fix_kinder)

        ebitda = [db[i] - fix_kopf.werte[i] for i in range(len(db))]
        z.append(Zeile('ebitda', 'EBITDA', 'geld', ebitda))

        afa = self._reihe(self.guv, GUV_AFA)
        if any(afa):
            zl = Zeile('afa', self._titel(self.guv, GUV_AFA, 'Abschreibungen'),
                       'geld', afa, ebene=1)
            zl.vorzeichen = -1
            z.append(zl)
        ebit = [ebitda[i] - afa[i] for i in range(len(ebitda))]
        z.append(Zeile('ebit', 'EBIT', 'geld', ebit))

        zins = [self._reihe(self.guv, GUV_ZINSAUFWAND)[i]
                - self._reihe(self.guv, GUV_ZINSERTRAG)[i] for i in range(len(ebit))]
        if any(zins):
            zl = Zeile('zins', 'Zinsergebnis', 'geld', zins, ebene=1)
            zl.vorzeichen = -1
            z.append(zl)
        ebt = [ebit[i] - zins[i] for i in range(len(ebit))]
        z.append(Zeile('ebt', 'EBT', 'geld', ebt))
        z.append(Zeile('ebt_marge', 'EBT-Marge', 'prozent',
                       [ebt[i] / umsatz_kopf.werte[i] * 100 if umsatz_kopf.werte[i]
                        else None for i in range(len(ebt))], ebene=2))
        self.bloecke.append(('ergebnis', 'Ergebnisrechnung', z))

    def _operativ(self):
        z = []
        menge = self._reihe(self.men, MEN['menge'], _w)
        kap = self._reihe(self.men, MEN['kapazitaet'], _w)
        umsatz = self._holen('umsatz')
        db = self._holen('db')

        z.append(Zeile('menge', self.einheit_plural, 'zahl', menge, ebene=1))
        if any(w is not None for w in kap):
            z.append(Zeile('kapazitaet', 'Kapazität', 'zahl', kap, ebene=2))
            z.append(Zeile('auslastung', 'Auslastung', 'prozent',
                           [_teile(menge[i], kap[i]) and menge[i] / kap[i] * 100
                            if kap[i] else None for i in range(len(menge))], ebene=2))
        z.append(Zeile('erloes_je', f'Ø Erlös je {self.einheit}', 'geld',
                       [_teile(umsatz[i], menge[i]) for i in range(len(menge))], ebene=2))
        # Wie im Bericht ohne die sonstigen betrieblichen Ertraege: die haengen
        # nicht an der Leistungsmenge und gehoeren deshalb nicht in den
        # Deckungsbeitrag je Einheit.
        db_menge = [self._holen('gesamtleistung')[i] - self._holen('sonstige_ertrag')[i]
                    - self._holen('variabel')[i] for i in range(len(menge))]
        z.append(Zeile('db_je', f'Deckungsbeitrag je {self.einheit}', 'geld',
                       [_teile(db_menge[i], menge[i]) for i in range(len(menge))], ebene=2))

        for schluessel, ersatz in [('koepfe', 'Mitarbeitende'), ('vzae', 'Vollzeitäquivalente'),
                                   ('krankenstand', 'Krankenstand'), ('austritte', 'Austritte'),
                                   ('aktive_kunden', 'Aktive Kunden'),
                                   ('neukunden', 'Neukunden je Monat'),
                                   ('auftragsbestand', 'Auftragsbestand'),
                                   ('reklamationen', 'Reklamationen')]:
            zeile = MEN[schluessel]
            werte = self._reihe(self.men, zeile, _w)
            if any(w is not None for w in werte):
                art = 'geld' if schluessel == 'auftragsbestand' else (
                    'prozent' if schluessel == 'krankenstand' else 'zahl')
                z.append(Zeile(schluessel, self._titel(self.men, zeile, ersatz),
                               art, werte, ebene=1))

        for zeile in EIGENE_ZEILEN:
            titel = self.men.cell(zeile, 1).value
            werte = self._reihe(self.men, zeile, _w)
            if titel and not str(titel).startswith('Eigene Kennzahl') \
                    and any(w is not None for w in werte):
                z.append(Zeile(f'eigen{zeile}', str(titel).strip(), 'zahl', werte, ebene=1))

        self.bloecke.append(('operativ', 'Menge und Betrieb', z))

    def _liquiditaet(self):
        z = []
        umsatz = self._holen('umsatz')
        liquide = self._reihe(self.liq, LIQ['liquide'], _w)
        ford = self._reihe(self.liq, LIQ['forderungen'], _w)
        verb = self._reihe(self.liq, LIQ['kurzfr_verb'], _w)
        ek = self._reihe(self.liq, LIQ['eigenkapital'], _w)
        bs = self._reihe(self.liq, LIQ['bilanzsumme'], _w)

        if any(w is not None for w in liquide):
            z.append(Zeile('liquide', self._titel(self.liq, LIQ['liquide'], 'Liquide Mittel'),
                           'geld', liquide, ebene=1))
        if any(w is not None for w in ford):
            z.append(Zeile('forderungen', 'Offene Forderungen', 'geld', ford, ebene=2))
            z.append(Zeile('dso', 'Debitorenlaufzeit', 'tage',
                           [ford[i] / umsatz[i] * 30 if ford[i] is not None and umsatz[i]
                            else None for i in range(len(ford))], ebene=2))
        if any(w is not None for w in verb):
            z.append(Zeile('kurzfr_verb', 'Kurzfristige Verbindlichkeiten', 'geld',
                           verb, ebene=2, vorzeichen=-1))
        if any(w is not None for w in ek) and any(w is not None for w in bs):
            z.append(Zeile('eigenkapital', 'Eigenkapital', 'geld', ek, ebene=1))
            z.append(Zeile('ek_quote', 'Eigenkapitalquote', 'prozent',
                           [ek[i] / bs[i] * 100 if ek[i] is not None and bs[i]
                            else None for i in range(len(ek))], ebene=2))
        self.bloecke.append(('liquiditaet', 'Liquidität und Bilanz', z))

    # ------------------------------------------------------------------ Ziele
    def _ziele_zuordnen(self):
        """Zielwerte stehen mit Beschriftung in Blatt 5. Zugeordnet wird ueber
        den Text, damit umbenannte Zeilen nicht ins Leere laufen."""
        nach_titel = {}
        for _, _, zeilen in self.bloecke:
            for zl in zeilen:
                nach_titel[zl.titel.strip().lower()] = zl

        fest = {'umsatzerlöse je monat': 'umsatz', 'ebt-marge': 'ebt_marge',
                'deckungsbeitragsquote': 'db_quote',
                'liquide mittel (mindestbestand)': 'liquide',
                'eigenkapitalquote': 'ek_quote', 'debitorenlaufzeit (dso)': 'dso',
                'auslastung': 'auslastung', 'aktive kunden / objekte': 'aktive_kunden',
                'neukunden je monat': 'neukunden'}
        nach_schluessel = {}
        for _, _, zeilen in self.bloecke:
            for zl in zeilen:
                nach_schluessel[zl.schluessel] = zl

        self.ohne_zuordnung = []
        for zeile in ZIEL_ZEILEN:
            titel = self.ziel.cell(zeile, 1).value
            wert = _w(self.ziel.cell(zeile, 2).value)
            richtung = self.ziel.cell(zeile, 3).value or 'höher ist besser'
            if not titel or wert is None:
                continue
            schluessel = fest.get(str(titel).strip().lower())
            zl = nach_schluessel.get(schluessel) if schluessel else None
            if zl is None:
                zl = nach_titel.get(str(titel).strip().lower())
            if zl is None:
                self.ohne_zuordnung.append(str(titel).strip())
                continue
            zl.ziel = wert
            zl.richtung = 'niedriger' if 'niedriger' in str(richtung) else 'hoeher'

    # --------------------------------------------------------------- Zugriff
    def _holen(self, schluessel):
        for _, _, zeilen in self.bloecke:
            for zl in zeilen:
                if zl.schluessel == schluessel:
                    return zl.werte
        return [0.0] * len(self.monate)

    def zeile(self, schluessel):
        for _, _, zeilen in self.bloecke:
            for zl in zeilen:
                if zl.schluessel == schluessel:
                    return zl
        return None

    def ampel(self, zl, i):
        """Gruen ab Ziel, Gelb bis zehn Prozent daneben, sonst Rot."""
        if zl.ziel is None or i >= len(zl.werte) or zl.werte[i] is None:
            return None
        ist, ziel = zl.werte[i], zl.ziel
        if not ziel:
            return None
        erfuellung = ist / ziel if zl.richtung == 'hoeher' else ziel / ist if ist else 0
        if (zl.richtung == 'hoeher' and ist >= ziel) or \
           (zl.richtung == 'niedriger' and ist <= ziel):
            return ('gruen', erfuellung)
        return ('gelb' if erfuellung >= 0.9 else 'rot', erfuellung)


if __name__ == '__main__':
    m = Matrix(sys.argv[1])
    print(m.firma, '|', m.berichtsmonat, m.jahr, '|', len(m.monate), 'Monate')
    for kennung, titel, zeilen in m.bloecke:
        print('==', titel)
        for zl in zeilen:
            w = zl.werte[m.aktiv]
            print(f'   {"  " * zl.ebene}{zl.titel:44s} {w if w is None else round(w, 2)}'
                  f'{"   Ziel " + str(zl.ziel) if zl.ziel is not None else ""}')
    if m.ohne_zuordnung:
        print('Zielzeilen ohne passende Kennzahl:', m.ohne_zuordnung)
