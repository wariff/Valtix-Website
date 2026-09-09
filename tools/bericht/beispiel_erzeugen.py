#!/usr/bin/env python3
"""Fuellt die leere Eingabevorlage mit Testzahlen.

Juli stammt eins zu eins aus dem bestehenden Word-Bericht der Lueftungstechnik
GmbH. Januar bis Juni sind so aufgeteilt, dass die dort veroeffentlichten
Monatswerte (Umsatz, Deckungsbeitrag, Fixkosten, EBT, Auslastung) exakt
herauskommen; die Aufteilung auf die einzelnen Kostenarten ist synthetisch.
Der Datensatz dient allein dazu, den Generator gegen bekannte Ergebnisse zu
pruefen, er bildet keinen echten Mandanten ab.
"""
import os, shutil, warnings
import openpyxl
warnings.filterwarnings('ignore')

HIER = os.path.dirname(os.path.abspath(__file__))
VORLAGE = os.path.join(HIER, 'VALTIX_Eingabevorlage.xlsx')
ZIEL    = os.path.join(HIER, 'beispiel_lueftungstechnik.xlsx')

# ---- Juli, direkt aus dem Bericht ----
JULI = dict(
    erloese_haupt=228196.36, erloese_material=136861.61, sonstige_ertraege=2345.60,
    material=116346.42, fremdleistung=17142.83, fahrzeug_var=9414.51,
    personal=142062.35, miete=11737.40, fuhrpark_fix=11124.35, marketing=2987.60,
    it=2340.75, versicherung=5476.30, sonstiges_fix=5876.90,
    afa=6680.90, zins=2567.45,
    menge=2810.30, kapazitaet=3444.0,
)
# Veroeffentlichte Monatswerte Januar bis Juni
MONATE = [
    ("Januar",   267251, 168797, 170807,  -4855, 72.6),
    ("Februar",  275718, 174736, 171180,    757, 74.9),
    ("März",     302304, 191674, 176885,  12038, 77.8),
    ("April",    299566, 190532, 177900,   9927, 76.3),
    ("Mai",      309788, 195379, 185697,   7024, 73.5),
    ("Juni",     333812.16, 203703.26, 182571.00, 18519.51, 79.2),
]
# Anteile der Kostenarten, abgeleitet aus dem Juli
VAR_ANTEIL = {'material': .8143, 'fremdleistung': .1199, 'fahrzeug_var': .0658}
FIX_ANTEIL = {'personal': .7822, 'miete': .0646, 'fuhrpark_fix': .0612,
              'marketing': .0164, 'it': .0129, 'versicherung': .0302,
              'sonstiges_fix': .0325}
SPALTE = {m: chr(ord('B')+i) for i, m in enumerate(
    ["Januar","Februar","März","April","Mai","Juni","Juli","August",
     "September","Oktober","November","Dezember"])}

shutil.copy(VORLAGE, ZIEL)
wb = openpyxl.load_workbook(ZIEL)

s = wb['1 Stammdaten']
for zeile, wert in [(5,'Lüftungstechnik GmbH'), (6,'GmbH'), (7,'Lüftungs- und Klimatechnik'),
                    (8,2026), (9,'Juli'), (10,'Sharif Ibrahim'), (11,'info@valtixfm.de'),
                    (14,'Monteurstunde'), (15,'Monteurstunden')]:
    s.cell(row=zeile, column=2, value=wert)

guv, men, liq = wb['2 GuV'], wb['3 Mengen & Operativ'], wb['4 Liquidität & Bilanz']

def setz(ws, zeile, monat, wert):
    ws[f'{SPALTE[monat]}{zeile}'] = round(wert, 2)

# Januar bis Juni aus den veroeffentlichten Aggregaten zurueckrechnen
for monat, umsatz, db, fix_inkl_afa, ebt, ausl in MONATE:
    var = umsatz - db                     # Gesamtleistung hier ohne sonstige Ertraege
    setz(guv, 6, monat, umsatz*0.625)     # Lohnerloese
    setz(guv, 9, monat, umsatz*0.375)     # Materialerloese
    for feld, zeile in (('material',15), ('fremdleistung',16)):
        setz(guv, zeile, monat, var*VAR_ANTEIL[feld])
    setz(guv, 21, monat, var*VAR_ANTEIL['fahrzeug_var'])
    afa = JULI['afa']
    fix = fix_inkl_afa - afa
    for feld, zeile in (('personal',27), ('miete',28), ('fuhrpark_fix',30),
                        ('marketing',34), ('it',33), ('versicherung',32),
                        ('sonstiges_fix',36)):
        setz(guv, zeile, monat, fix*FIX_ANTEIL[feld])
    setz(guv, 41, monat, afa)
    setz(guv, 43, monat, db - fix - afa - ebt)   # Zinsaufwand als Restgroesse
    menge = ausl/100 * JULI['kapazitaet']
    setz(men, 6, monat, menge)
    setz(men, 7, monat, JULI['kapazitaet'])

# Juli exakt
setz(guv,  6, 'Juli', JULI['erloese_haupt'])
setz(guv,  9, 'Juli', JULI['erloese_material'])
setz(guv, 11, 'Juli', JULI['sonstige_ertraege'])
setz(guv, 15, 'Juli', JULI['material'])
setz(guv, 16, 'Juli', JULI['fremdleistung'])
setz(guv, 21, 'Juli', JULI['fahrzeug_var'])
for feld, zeile in (('personal',27), ('miete',28), ('fuhrpark_fix',30),
                    ('marketing',34), ('it',33), ('versicherung',32),
                    ('sonstiges_fix',36)):
    setz(guv, zeile, 'Juli', JULI[feld])
setz(guv, 41, 'Juli', JULI['afa'])
setz(guv, 43, 'Juli', JULI['zins'])
setz(men,  6, 'Juli', JULI['menge'])
setz(men,  7, 'Juli', JULI['kapazitaet'])

# Operative Kennzahlen und Liquiditaet
for monat, aktiv, neu, bestand in [("Januar",41,7,512000),("Februar",43,5,534800),
        ("März",45,9,571400),("April",46,6,598200),("Mai",47,8,616900),
        ("Juni",48,7,645230),("Juli",49,6,718561)]:
    setz(men, 21, monat, aktiv); setz(men, 22, monat, neu); setz(men, 23, monat, bestand)
    setz(men, 13, monat, 24); setz(men, 14, monat, 22.5)

for monat, kasse, ford, kurzfr, ek, bs in [
        ("Januar",148300,392000,498000,742000,2098000),
        ("Februar",141200,404500,505000,748000,2104000),
        ("März",137800,441000,516000,760000,2118000),
        ("April",132400,455000,523000,770000,2130000),
        ("Mai",128900,470000,531000,777000,2142000),
        ("Juni",125300,512000,544000,795000,2178000),
        ("Juli",121560.45,584092.75,561000,801400,2238000)]:
    setz(liq, 6, monat, kasse); setz(liq, 7, monat, ford); setz(liq, 8, monat, kurzfr)
    setz(liq, 13, monat, ek);   setz(liq, 14, monat, bs)

ziel = wb['5 Zielwerte']
for zeile, wert in [(5,335000),(6,8.0),(7,40.0),(8,150000),(9,35.0),
                    (10,40),(11,78.0),(12,40),(13,8)]:
    ziel.cell(row=zeile, column=2, value=wert)

wb.save(ZIEL)
print('geschrieben:', ZIEL)
