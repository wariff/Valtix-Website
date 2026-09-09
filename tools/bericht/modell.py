#!/usr/bin/env python3
"""Liest die Valtix-Eingabevorlage und rechnet das Modell durch.

Bewusst werden nur die Eingabefelder gelesen. Die Summen der Vorlage sind
Excel-Formeln, deren Ergebnis in einer nie geoeffneten Datei nicht gespeichert
ist. Alles Abgeleitete wird hier gerechnet, damit der Bericht auch aus einer
Datei entsteht, die der Mandant nur ausgefuellt und gespeichert hat.
"""
import warnings
import openpyxl
warnings.filterwarnings('ignore')

MONATE = ["Januar","Februar","März","April","Mai","Juni","Juli","August",
          "September","Oktober","November","Dezember"]

# Zeilennummern der Eingabefelder je Blatt
GUV_ERLOESE   = {'Erlöse Hauptleistung':6, 'Erlöse Nebenleistung 1':7,
                 'Erlöse Nebenleistung 2':8, 'Warenverkauf / Material':9}
GUV_SONSTIGE  = 11
GUV_VARIABEL  = {'Wareneinsatz / Materialaufwand':15, 'Fremdleistungen':16,
                 'Leistungsabhängige Löhne':17, 'Energie mengenabhängig':18,
                 'Provisionen / Plattformgebühren':19, 'Verpackung / Transport':20,
                 'Sonstige variable Kosten':21}
GUV_FIX       = {'Personalaufwand':27, 'Miete & Nebenkosten':28,
                 'Energie (Grundgebühren)':29, 'Fahrzeuge / Fuhrpark':30,
                 'Wartung & Instandhaltung':31, 'Versicherungen & Beiträge':32,
                 'IT, Software & Kommunikation':33, 'Marketing & Vertrieb':34,
                 'Beratung':35, 'Sonstige betriebliche Aufwendungen':36}
GUV_AFA, GUV_ZINSAUFWAND, GUV_ZINSERTRAG = 41, 43, 44

MEN = {'menge':6, 'kapazitaet':7, 'koepfe':13, 'vzae':14, 'bezahlte_stunden':15,
       'krankenstand':17, 'austritte':18, 'aktive_kunden':21, 'neukunden':22,
       'auftragsbestand':23, 'reklamationen':24}
LIQ = {'liquide':6, 'forderungen':7, 'kurzfr_verb':8,
       'eigenkapital':13, 'bilanzsumme':14, 'finanzverb':16, 'anlagevermoegen':17}

ZIELE = [  # Zeile, Bezeichnung, Richtung, Einheit
    (5,  'Umsatzerlöse',            'hoch', 'eur'),
    (6,  'EBT-Marge',               'hoch', 'proz'),
    (7,  'Deckungsbeitragsquote',   'hoch', 'proz'),
    (8,  'Liquide Mittel',          'hoch', 'eur'),
    (9,  'Eigenkapitalquote',       'hoch', 'proz'),
    (10, 'Debitorenlaufzeit (DSO)', 'tief', 'tage'),
    (11, 'Auslastung',              'hoch', 'proz'),
    (12, 'Aktive Kunden',           'hoch', 'zahl'),
    (13, 'Neukunden je Monat',      'hoch', 'zahl'),
]


def _z(wert):
    """Leere Zellen und Text als 0 behandeln, Zahlen als float."""
    if wert is None or isinstance(wert, str):
        return 0.0
    return float(wert)


class Bericht:
    def __init__(self, pfad):
        wb = openpyxl.load_workbook(pfad, data_only=True)
        self.wb = wb
        self.stamm = self._stammdaten(wb['1 Stammdaten'])
        self.monate = self._monate(wb)
        self.ziele = self._ziele(wb['5 Zielwerte'])
        if not self.monate:
            raise SystemExit('Keine Monate mit Umsatz gefunden. Ist die Datei ausgefüllt?')
        self.aktuell = self.stamm['berichtsmonat'] or self.monate[-1]['monat']
        self.idx = [m['monat'] for m in self.monate].index(self.aktuell)

    # ---------- Einlesen ----------
    def _stammdaten(self, ws):
        holen = lambda z: ws.cell(row=z, column=2).value
        return dict(firma=holen(5) or 'Unbekannt', rechtsform=holen(6) or '',
                    branche=holen(7) or '', jahr=holen(8) or '',
                    berichtsmonat=holen(9), ansprechpartner=holen(10) or '',
                    email=holen(11) or '', einheit=holen(14) or 'Leistungseinheit',
                    einheit_plural=holen(15) or 'Leistungseinheiten')

    def _monate(self, wb):
        guv, men, liq = wb['2 GuV'], wb['3 Mengen & Operativ'], wb['4 Liquidität & Bilanz']
        raus = []
        for i, name in enumerate(MONATE):
            sp = 2 + i
            erloese = {k: _z(guv.cell(row=z, column=sp).value) for k, z in GUV_ERLOESE.items()}
            umsatz = sum(erloese.values())
            if umsatz <= 0:
                continue
            var = {k: _z(guv.cell(row=z, column=sp).value) for k, z in GUV_VARIABEL.items()}
            fix = {k: _z(guv.cell(row=z, column=sp).value) for k, z in GUV_FIX.items()}
            sonst = _z(guv.cell(row=GUV_SONSTIGE, column=sp).value)
            afa   = _z(guv.cell(row=GUV_AFA, column=sp).value)
            zins  = _z(guv.cell(row=GUV_ZINSAUFWAND, column=sp).value) \
                    - _z(guv.cell(row=GUV_ZINSERTRAG, column=sp).value)

            gesamtleistung = umsatz + sonst
            var_summe = sum(var.values())
            db = gesamtleistung - var_summe
            fix_summe = sum(fix.values())
            ebitda = db - fix_summe
            ebit = ebitda - afa
            ebt = ebit - zins

            op  = {k: _z(men.cell(row=z, column=sp).value) for k, z in MEN.items()}
            fin = {k: _z(liq.cell(row=z, column=sp).value) for k, z in LIQ.items()}

            raus.append(dict(
                monat=name, erloese=erloese, umsatz=umsatz, sonstige=sonst,
                gesamtleistung=gesamtleistung, variabel=var, var_summe=var_summe,
                db=db, db_quote=db / gesamtleistung * 100 if gesamtleistung else 0,
                fix=fix, fix_summe=fix_summe, afa=afa, fix_inkl_afa=fix_summe + afa,
                ebitda=ebitda, ebit=ebit, zins=zins, ebt=ebt,
                ebt_marge=ebt / umsatz * 100 if umsatz else 0,
                op=op, fin=fin,
                auslastung=op['menge'] / op['kapazitaet'] * 100 if op['kapazitaet'] else 0,
                erloes_je_einheit=umsatz / op['menge'] if op['menge'] else 0,
                db_je_einheit=(umsatz - var_summe) / op['menge'] if op['menge'] else 0,
                liq_1=fin['liquide'] / fin['kurzfr_verb'] * 100 if fin['kurzfr_verb'] else 0,
                dso=fin['forderungen'] / umsatz * 30 if umsatz else 0,
                ek_quote=fin['eigenkapital'] / fin['bilanzsumme'] * 100 if fin['bilanzsumme'] else 0,
            ))
        return raus

    def _ziele(self, ws):
        raus = []
        for zeile, name, richtung, einheit in ZIELE:
            wert = ws.cell(row=zeile, column=2).value
            if wert is None or isinstance(wert, str):
                continue
            raus.append(dict(zeile=zeile, name=name, richtung=richtung,
                             einheit=einheit, ziel=float(wert)))
        return raus

    # ---------- Auswertungen ----------
    @property
    def m(self):
        return self.monate[self.idx]

    @property
    def vormonat(self):
        return self.monate[self.idx - 1] if self.idx > 0 else None

    def ampel(self, ist, ziel, richtung):
        """gruen = Ziel erreicht, gelb = bis 10 Prozent verfehlt, sonst rot."""
        if ziel == 0:
            return 'grau'
        if richtung == 'hoch':
            if ist >= ziel: return 'gruen'
            return 'gelb' if ist >= ziel * 0.9 else 'rot'
        if ist <= ziel: return 'gruen'
        return 'gelb' if ist <= ziel * 1.1 else 'rot'

    def ist_wert(self, name):
        m = self.m
        return {'Umsatzerlöse': m['umsatz'], 'EBT-Marge': m['ebt_marge'],
                'Deckungsbeitragsquote': m['db_quote'],
                'Liquide Mittel': m['fin']['liquide'],
                'Eigenkapitalquote': m['ek_quote'],
                'Debitorenlaufzeit (DSO)': m['dso'],
                'Auslastung': m['auslastung'],
                'Aktive Kunden': m['op']['aktive_kunden'],
                'Neukunden je Monat': m['op']['neukunden']}.get(name, 0)

    def break_even(self):
        """Fuer die Gewinnschwelle zaehlt nur der mengenabhaengige Teil.
        Sonstige betriebliche Ertraege haengen nicht an der Leistungsmenge und
        bleiben deshalb im Deckungsbeitrag je Einheit aussen vor."""
        m = self.m
        if not m['op']['menge']:
            return None
        db_je = (m['umsatz'] - m['var_summe']) / m['op']['menge']
        if db_je <= 0:
            return None
        schwelle = m['fix_inkl_afa'] / db_je
        cash = m['fix_summe'] / db_je
        umsatz_je = m['erloes_je_einheit']
        return dict(
            db_je_einheit=db_je, umsatz_je_einheit=umsatz_je,
            schwelle=schwelle, schwelle_umsatz=schwelle * umsatz_je,
            cash_schwelle=cash, cash_umsatz=cash * umsatz_je,
            ist=m['op']['menge'], abstand=m['op']['menge'] - schwelle,
            abstand_proz=(m['op']['menge'] - schwelle) / m['op']['menge'] * 100
                          if m['op']['menge'] else 0,
            puffer_umsatz=(m['op']['menge'] - schwelle) * umsatz_je,
            kapazitaet=m['op']['kapazitaet'],
            kapazitaet_umsatz=m['op']['kapazitaet'] * umsatz_je,
        )


if __name__ == '__main__':
    import sys
    b = Bericht(sys.argv[1] if len(sys.argv) > 1 else 'beispiel_lueftungstechnik.xlsx')
    m = b.m
    print(f"{b.stamm['firma']} | {b.aktuell} {b.stamm['jahr']} | {len(b.monate)} Monate erfasst")
    for k in ('umsatz','gesamtleistung','var_summe','db','db_quote','fix_summe',
              'fix_inkl_afa','ebitda','afa','ebit','zins','ebt','ebt_marge',
              'auslastung','liq_1','dso','ek_quote'):
        print(f'  {k:16} {m[k]:>14,.2f}'.replace(',', '.'))
    be = b.break_even()
    print('  Break-even      ', f"{be['schwelle']:,.2f} Std / {be['schwelle_umsatz']:,.2f} EUR".replace(',', '.'))
    print('  DB je Einheit   ', f"{be['db_je_einheit']:,.2f}".replace(',', '.'))
    print('  Sicherheit      ', f"{be['abstand']:,.2f} Std ({be['abstand_proz']:.1f} %)".replace(',', '.'))
    print('  Umsatzpuffer    ', f"{be['puffer_umsatz']:,.2f}".replace(',', '.'))
